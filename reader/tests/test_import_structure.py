import tempfile
import unittest
from pathlib import Path

from core.parser import docx_parser, epub_parser, pdf_parser, txt_parser
from core.parser.structure import attach_blocks

PROSE = 'A sufficiently long paragraph retains its own shape and can be narrated naturally. ' * 12


class ImportStructureTests(unittest.TestCase):
    def test_html_inline_markup_and_source_indentation_do_not_split_paragraph(self):
        parser = epub_parser._HTMLLineExtractor()
        parser.feed('<h1>Chapter 1</h1><p>Hello <em>dear</em>\nreader.</p><h2>A detail</h2><p>Second paragraph.</p>')
        self.assertEqual(parser.get_lines(), ['Chapter 1', 'Hello dear reader.', 'A detail', 'Second paragraph.'])
        self.assertEqual(parser.kinds['A detail'], 'subheading')

    def test_epub_split_retains_heading_once(self):
        _, sections = epub_parser._split_document(['Chapter 1', PROSE, 'A second paragraph.'])
        self.assertTrue(sections[0]['content'].startswith('Chapter 1\n\n'))
        self.assertEqual(sections[0]['content'].count('Chapter 1'), 1)

    def test_docx_keeps_nested_headings_in_chapter(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'book.docx'
            doc = Document()
            doc.add_heading('Chapter 1', level=1)
            doc.add_paragraph(PROSE)
            doc.add_heading('A quiet moment', level=2)
            doc.add_paragraph('The next paragraph.')
            doc.save(path)
            chapters = docx_parser.parse(path)['chapters']
        self.assertEqual(len(chapters), 1)
        self.assertEqual([b['kind'] for b in chapters[0]['blocks']], ['heading', 'paragraph', 'subheading', 'paragraph'])
        self.assertEqual(chapters[0]['content'].count('Chapter 1'), 1)

    def test_txt_preserves_paragraphs_and_real_heading_only(self):
        chapter = txt_parser.parse_text('Chapter 1\n' + PROSE + '\n\nSecond paragraph.', title='Metadata title')['chapters'][0]
        self.assertEqual(chapter['blocks'][0]['kind'], 'heading')
        self.assertEqual(len(chapter['blocks']), 3)
        self.assertNotIn('Metadata title', chapter['content'])

    def test_pdf_retains_title_and_detected_subheading(self):
        blocks = [{'text': 'Chapter 1', 'size': 20}, {'text': PROSE, 'size': 12},
                  {'text': 'A quiet moment', 'size': 15}, {'text': PROSE, 'size': 12}]
        chapter = pdf_parser._split_chapters(blocks, 'Metadata title')[0]
        self.assertEqual([b['kind'] for b in chapter['blocks']], ['heading', 'paragraph', 'subheading', 'paragraph'])

    def test_fallback_blocks_have_optional_controls_without_invented_title(self):
        chapter = attach_blocks([{'title': 'Fallback', 'content': 'One.\n\nTwo.'}])[0]
        self.assertEqual([b['text'] for b in chapter['blocks']], ['One.', 'Two.'])
        self.assertTrue(all(b['speed'] is None and b['pause_ms'] is None for b in chapter['blocks']))

    def test_epub_file_import_keeps_semantic_blocks(self):
        from ebooklib import epub
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'book.epub'
            book = epub.EpubBook()
            book.set_identifier('structure-test')
            book.set_title('Metadata title')
            book.set_language('en')
            section = epub.EpubHtml(title='Chapter 1', file_name='chapter.xhtml', lang='en')
            section.content = '<h1>Chapter 1</h1><p>' + PROSE + '</p><h2>A quiet moment</h2><p>Second paragraph.</p>'
            book.add_item(section)
            book.toc = (section,)
            book.spine = [section]
            book.add_item(epub.EpubNcx())
            book.add_item(epub.EpubNav())
            epub.write_epub(str(path), book)
            chapter = epub_parser.parse(path)['chapters'][0]
        self.assertEqual([b['kind'] for b in chapter['blocks']], ['heading', 'paragraph', 'subheading', 'paragraph'])
        self.assertEqual(chapter['content'], '\n\n'.join(b['text'] for b in chapter['blocks']))

    def test_html_br_remains_inside_one_paragraph(self):
        parser = epub_parser._HTMLLineExtractor()
        parser.feed('<p>First line.<br/>Second line.</p><p>New paragraph.</p>')
        lines = parser.get_lines()
        self.assertEqual(lines, ['First line.\nSecond line.', 'New paragraph.'])
        chapter = attach_blocks([{'title': 'Metadata', 'content': '\n\n'.join(lines)}])[0]
        self.assertEqual(len(chapter['blocks']), 2)

    def test_docx_subtitle_is_not_a_chapter_boundary(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'book.docx'
            doc = Document()
            doc.add_heading('Chapter 1', level=1)
            doc.add_paragraph('A descriptive subtitle', style='Subtitle')
            doc.add_paragraph(PROSE)
            doc.save(path)
            chapters = docx_parser.parse(path)['chapters']
        self.assertEqual(len(chapters), 1)
        self.assertEqual(chapters[0]['blocks'][1]['kind'], 'subheading')

    def test_pdf_continues_sentence_across_page_with_nontext_first_block(self):
        blocks = [{'text': 'The sentence continues', 'size': 12, 'page': 0},
                  {'text': 'on another page.', 'size': 12, 'page': 1, 'paragraph_start': True},
                  {'text': 'A new paragraph.', 'size': 12, 'page': 1, 'paragraph_start': True}]
        self.assertEqual(pdf_parser._merge_pdf_blocks(blocks),
                         'The sentence continues on another page.\n\nA new paragraph.')

    def test_txt_hard_wrapping_does_not_create_extra_blocks(self):
        chapter = txt_parser.parse_text('Chapter 1\n' + PROSE[:PROSE.index(' ', 70)] + '\n' + PROSE[PROSE.index(' ', 70) + 1:] + '\n\nFinal paragraph.')['chapters'][0]
        self.assertEqual(len(chapter['blocks']), 3)
        self.assertEqual(' '.join(chapter['blocks'][1]['text'].split()), ' '.join(PROSE.split()))

    def test_parser_blocks_validate_and_enrich_without_losing_headings(self):
        from core import text_editor
        chapter = txt_parser.parse_text('Chapter 1\n' + PROSE)['chapters'][0]
        validated = text_editor.validate_blocks(chapter['blocks'])
        self.assertEqual(text_editor.content_of(validated), chapter['content'])
        segments = text_editor.enrich_blocks(validated, [], None, True, None)
        self.assertEqual(segments[0]['block_kind'], 'heading')
        self.assertEqual(segments[0]['enriched_text'], 'Chapter 1.')
        self.assertTrue(any(s['block_kind'] == 'paragraph' for s in segments))

    def test_epub_toc_only_title_and_repeated_text_keep_own_kinds(self):
        from ebooklib import epub
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'book.epub'
            book = epub.EpubBook()
            book.set_identifier('repeated-structure-test')
            book.set_title('Metadata title')
            book.set_language('en')
            sections = []
            for i, (title, html) in enumerate([
                ('Only in TOC', '<p>' + PROSE + '</p>'),
                ('Echo', '<h1>Echo</h1><p>Echo</p><p>' + PROSE + '</p>'),
                ('Other title', '<p>Echo</p><h2>Echo</h2><p>' + PROSE + '</p>'),
            ]):
                section = epub.EpubHtml(title=title, file_name=f'chapter{i}.xhtml', lang='en')
                section.content = html
                book.add_item(section)
                sections.append(section)
            book.toc = tuple(sections)
            book.spine = sections
            book.add_item(epub.EpubNcx())
            book.add_item(epub.EpubNav())
            epub.write_epub(str(path), book)
            chapters = epub_parser.parse(path)['chapters']
        self.assertEqual(chapters[0]['blocks'][0]['text'], 'Only in TOC')
        self.assertEqual(chapters[0]['blocks'][0]['kind'], 'heading')
        self.assertEqual([b['kind'] for b in chapters[1]['blocks']], ['heading', 'paragraph', 'paragraph'])
        self.assertEqual(chapters[1]['content'].count('Echo'), 2)
        self.assertEqual([b['kind'] for b in chapters[2]['blocks']], ['heading', 'paragraph', 'subheading', 'paragraph'])

    def test_docx_repeated_heading_wording_keeps_paragraph_occurrence(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'book.docx'
            doc = Document()
            doc.add_heading('Echo', level=1)
            doc.add_paragraph('Echo')
            doc.add_heading('Echo', level=2)
            doc.add_paragraph(PROSE)
            doc.save(path)
            chapter = docx_parser.parse(path)['chapters'][0]
        self.assertEqual([b['kind'] for b in chapter['blocks']], ['heading', 'paragraph', 'subheading', 'paragraph'])
